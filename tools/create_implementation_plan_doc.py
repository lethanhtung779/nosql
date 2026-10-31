from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

OUT = Path(r"D:\nosql\doannosql\Ke_Hoach_Trien_Khai_He_Thong_Spa_Salon_NoSQL.docx")

doc = Document()
section = doc.sections[0]
section.page_width = Cm(21)
section.page_height = Cm(29.7)
section.top_margin = Cm(2.0)
section.bottom_margin = Cm(1.8)
section.left_margin = Cm(2.2)
section.right_margin = Cm(2.0)

styles = doc.styles
styles['Normal'].font.name = 'Arial'
styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
styles['Normal'].font.size = Pt(10.5)
styles['Normal'].paragraph_format.space_after = Pt(6)
styles['Normal'].paragraph_format.line_spacing = 1.15

for name, size, before, after in [('Title', 25, 0, 16), ('Heading 1', 16, 16, 8), ('Heading 2', 13, 12, 6), ('Heading 3', 11, 9, 4)]:
    s = styles[name]
    s.font.name = 'Arial'
    s._element.rPr.rFonts.set(qn('w:eastAsia'), 'Arial')
    s.font.size = Pt(size)
    s.font.bold = True
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.paragraph_format.space_before = Pt(before)
    s.paragraph_format.space_after = Pt(after)
    s.paragraph_format.keep_with_next = True

# Remove the decorative border inherited by some Word Title styles.
title_ppr = styles['Title']._element.get_or_add_pPr()
title_border = title_ppr.find(qn('w:pBdr'))
if title_border is not None:
    title_ppr.remove(title_border)

if 'Code' not in styles:
    code_style = styles.add_style('Code', WD_STYLE_TYPE.PARAGRAPH)
    code_style.font.name = 'Consolas'
    code_style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Consolas')
    code_style.font.size = Pt(9)
    code_style.paragraph_format.left_indent = Cm(0.7)
    code_style.paragraph_format.space_after = Pt(5)


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for margin, value in [('top', top), ('start', start), ('bottom', bottom), ('end', end)]:
        node = tcMar.find(qn(f'w:{margin}'))
        if node is None:
            node = OxmlElement(f'w:{margin}')
            tcMar.append(node)
        node.set(qn('w:w'), str(value))
        node.set(qn('w:type'), 'dxa')


def set_table_borders(table, color='D9D9D9', size='6'):
    tblPr = table._tbl.tblPr
    borders = tblPr.first_child_found_in('w:tblBorders')
    if borders is None:
        borders = OxmlElement('w:tblBorders')
        tblPr.append(borders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        tag = borders.find(qn(f'w:{edge}'))
        if tag is None:
            tag = OxmlElement(f'w:{edge}')
            borders.append(tag)
        tag.set(qn('w:val'), 'single')
        tag.set(qn('w:sz'), size)
        tag.set(qn('w:color'), color)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)


def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_table_borders(t)
    hdr = t.rows[0]
    repeat_header(hdr)
    for i, text in enumerate(headers):
        cell = hdr.cells[i]
        set_cell_shading(cell, '24435A')
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(str(text))
        r.bold = True
        r.font.color.rgb = RGBColor(255, 255, 255)
        r.font.size = Pt(9.5)
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            if ri % 2 == 1:
                set_cell_shading(cell, 'F3F6F8')
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 and len(headers) > 2 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            r.font.size = Pt(9.2)
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t


def bullet(text, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def number(text):
    p = doc.add_paragraph(style='List Number')
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p


def paragraph(text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def code(text):
    p = doc.add_paragraph(style='Code')
    p.add_run(text)
    return p


def page_break():
    doc.add_page_break()


# Cover
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(45)
r = p.add_run('ĐỒ ÁN CƠ SỞ DỮ LIỆU NOSQL')
r.bold = True; r.font.size = Pt(13); r.font.color.rgb = RGBColor(0, 0, 0)

p = doc.add_paragraph(style='Title')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(45)
p.add_run('KẾ HOẠCH TRIỂN KHAI HỆ THỐNG QUẢN LÝ LỊCH HẸN SPA SALON')

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(15)
r = p.add_run('Nền tảng MongoDB và Redis')
r.bold = True; r.font.size = Pt(15); r.font.color.rgb = RGBColor(0, 0, 0)

doc.add_paragraph().paragraph_format.space_before = Pt(90)
table(['Thành phần', 'Công nghệ'], [
    ['Frontend', 'Next.js và React'],
    ['Backend', 'NestJS và TypeScript'],
    ['CSDL NoSQL chính', 'MongoDB'],
    ['CSDL NoSQL hỗ trợ', 'Redis'],
    ['Môi trường', 'Docker Compose'],
], widths=[6, 10])

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(60)
r = p.add_run('Tháng 10 năm 2026')
r.font.size = Pt(11)

page_break()

doc.add_heading('Tóm tắt kế hoạch', level=1)
paragraph('Dự án xây dựng một hệ thống quản lý lịch hẹn Spa và Salon với MongoDB làm cơ sở dữ liệu nghiệp vụ chính và Redis làm cơ sở dữ liệu hỗ trợ thời gian thực. Hệ thống tập trung vào đặt lịch nhiều dịch vụ, chống trùng khung giờ, giữ chỗ tự hết hạn, quản lý quy trình phục vụ và tổng hợp báo cáo kinh doanh.')
paragraph('Kế hoạch triển khai được chia thành sáu tuần, bắt đầu từ phân tích yêu cầu và thiết kế dữ liệu, tiếp tục với hạ tầng, xác thực, nghiệp vụ đặt lịch, giao diện người dùng, quản trị, kiểm thử và hoàn thiện báo cáo. Kết quả cuối cùng là một ứng dụng có thể chạy bằng Docker Compose, có Swagger, dữ liệu mẫu, kiểm thử tự động và kịch bản trình diễn đầy đủ.')

doc.add_heading('Nội dung tài liệu', level=1)
for item in [
    '1 Mục tiêu và phạm vi dự án', '2 Kiến trúc tổng thể', '3 Thiết kế dữ liệu MongoDB',
    '4 Thiết kế dữ liệu Redis', '5 Luồng đặt lịch và chống tranh chấp', '6 Danh sách API',
    '7 Kế hoạch giao diện', '8 Kế hoạch triển khai theo giai đoạn', '9 Kế hoạch kiểm thử',
    '10 Tiêu chí nghiệm thu', '11 Rủi ro và biện pháp xử lý', '12 Kịch bản trình bày đồ án'
]:
    bullet(item)

doc.add_heading('1 Mục tiêu và phạm vi dự án', level=1)
doc.add_heading('1 1 Mục tiêu', level=2)
paragraph('Xây dựng hệ thống quản lý lịch hẹn và dịch vụ Spa Salon sử dụng ít nhất hai hệ quản trị cơ sở dữ liệu NoSQL. MongoDB chịu trách nhiệm lưu dữ liệu lâu dài; Redis chịu trách nhiệm giữ chỗ tạm thời, khóa chống tranh chấp và cache. Backend được tổ chức bằng NestJS, frontend dùng Next.js và toàn bộ hạ tầng được đóng gói bằng Docker Compose.')

doc.add_heading('1 2 Kết quả cần đạt', level=2)
for item in [
    'Khách hàng có thể đăng ký, đăng nhập, chọn chi nhánh, dịch vụ, chuyên viên và đặt lịch.',
    'Hai khách không thể giữ hoặc xác nhận cùng một khung giờ giao nhau của một chuyên viên.',
    'Redis tự động giải phóng lượt giữ chỗ sau 15 phút.',
    'Nhân viên theo dõi lịch làm việc và cập nhật trạng thái phục vụ.',
    'Quản trị viên quản lý danh mục và theo dõi báo cáo kinh doanh.',
    'Hệ thống có tài liệu API, kiểm thử, dữ liệu mẫu và hướng dẫn triển khai.'
]: bullet(item)

doc.add_heading('1 3 Phạm vi chức năng theo vai trò', level=2)
table(['Vai trò', 'Chức năng chính', 'Giới hạn truy cập'], [
    ['Khách hàng', 'Đăng ký, đăng nhập, tìm chi nhánh, chọn dịch vụ, giữ slot, xác nhận hoặc hủy lịch, đánh giá', 'Chỉ xem và thao tác dữ liệu của bản thân'],
    ['Nhân viên', 'Xem lịch cá nhân, xem dịch vụ và ghi chú, cập nhật trạng thái phục vụ', 'Chỉ xem lịch được phân công'],
    ['Quản trị viên', 'Quản lý chi nhánh, dịch vụ, nhân viên, lịch hẹn và báo cáo', 'Toàn quyền nghiệp vụ trong hệ thống'],
], widths=[3.2, 9.2, 4.4])

doc.add_heading('2 Kiến trúc tổng thể', level=1)
paragraph('Kiến trúc tách giao diện, API và hai lớp lưu trữ NoSQL. Next.js giao tiếp với NestJS bằng REST API và JWT. NestJS điều phối MongoDB và Redis theo đúng trách nhiệm của từng hệ cơ sở dữ liệu.')

table(['Tầng', 'Thành phần', 'Trách nhiệm'], [
    ['Giao diện', 'Next.js và React', 'Hiển thị dịch vụ, đặt lịch, dashboard nhân viên và quản trị'],
    ['API', 'NestJS', 'Xác thực, phân quyền, validation, xử lý nghiệp vụ và Swagger'],
    ['Dữ liệu chính', 'MongoDB', 'Lưu users, branches, services, appointments, reviews và staff calendars'],
    ['Dữ liệu thời gian thực', 'Redis', 'Giữ slot TTL, khóa tranh chấp và cache danh mục'],
    ['Hạ tầng', 'Docker Compose', 'Khởi động MongoDB, Redis và cấu hình môi trường thống nhất'],
], widths=[3, 4.3, 9.5])

doc.add_heading('2 1 Sơ đồ luồng thành phần', level=2)
table(['Người dùng', 'Frontend', 'Backend', 'Dữ liệu'], [
    ['Khách hàng Nhân viên Admin', 'Next.js cổng 3000', 'NestJS cổng 3001', 'MongoDB cổng 27017\nRedis cổng 6379'],
], widths=[4, 4, 4, 4.5])

doc.add_heading('2 2 Trách nhiệm của MongoDB', level=2)
for item in [
    'Lưu tài khoản, chi nhánh, dịch vụ, lịch hẹn, đánh giá và lịch nhân viên.',
    'Lưu snapshot dịch vụ, giá và nhân viên tại thời điểm đặt lịch.',
    'Truy vấn chi nhánh gần nhất bằng GeoJSON và chỉ mục 2dsphere.',
    'Tổng hợp doanh thu, dịch vụ phổ biến và hiệu suất bằng Aggregation Pipeline.',
    'Cập nhật staff calendars nguyên tử để kiểm tra xung đột lần cuối.'
]: bullet(item)

doc.add_heading('2 3 Trách nhiệm của Redis', level=2)
for item in [
    'Giữ khung giờ tạm thời trong 900 giây.',
    'Ngăn hai khách giữ các khoảng thời gian giao nhau của cùng nhân viên.',
    'Lưu chỉ mục slot theo nhân viên và ngày bằng Sorted Set.',
    'Cache danh sách chi nhánh và dịch vụ trong 300 giây.',
    'Mở rộng cho rate limiting và quản lý phiên khi cần.'
]: bullet(item)

doc.add_heading('3 Thiết kế dữ liệu MongoDB', level=1)
paragraph('Mô hình kết hợp Embedded Pattern cho dữ liệu thường xuyên đọc cùng nhau và Referenced Pattern cho dữ liệu độc lập, được dùng chung. Các trường quan trọng được gắn chỉ mục phù hợp với luồng truy vấn.')
table(['Collection', 'Dữ liệu chính', 'Chỉ mục quan trọng'], [
    ['users', 'Thông tin tài khoản, vai trò, chi nhánh và chuyên môn', 'phone unique; role và branch_id'],
    ['branches', 'Địa chỉ, giờ hoạt động và tọa độ GeoJSON', 'location 2dsphere'],
    ['services', 'Giá, thời lượng, quy trình và chi nhánh áp dụng', 'branch_ids và status'],
    ['appointments', 'Snapshot khách, nhân viên, dịch vụ, thời gian và tổng tiền', 'booking_code unique; staff và start_time'],
    ['staff_calendars', 'Mảng slot theo nhân viên và ngày', 'staff_id và date_key unique'],
    ['reviews', 'Điểm số, nhận xét và hình ảnh đánh giá', 'appointment_id unique; staff_id'],
], widths=[3.2, 8.1, 5.5])

doc.add_heading('3 1 Cấu trúc lịch hẹn', level=2)
code('appointments {')
code('  booking_code, branch_id,')
code('  customer: { user_id, name, phone },')
code('  staff: { staff_id, name },')
code('  time_slot: { start_time, end_time },')
code('  services_booked: [{ service_id, name, price, duration }],')
code('  total_amount, status, note, created_at, updated_at')
code('}')

doc.add_heading('3 2 Trạng thái lịch hẹn', level=2)
table(['Trạng thái', 'Ý nghĩa', 'Chuyển tiếp hợp lệ'], [
    ['CONFIRMED', 'Lịch đã được khách xác nhận', 'IN_PROGRESS hoặc CANCELLED'],
    ['IN_PROGRESS', 'Nhân viên đang thực hiện dịch vụ', 'COMPLETED'],
    ['COMPLETED', 'Dịch vụ đã hoàn thành', 'Cho phép khách đánh giá'],
    ['CANCELLED', 'Lịch đã bị hủy', 'Không chuyển tiếp'],
], widths=[3.5, 7.5, 5.8])

doc.add_heading('4 Thiết kế dữ liệu Redis', level=1)
table(['Nhóm khóa', 'Cấu trúc khóa', 'Dữ liệu và TTL'], [
    ['Giữ chỗ', 'slot-hold:{token}', 'JSON chứa người dùng, nhân viên, dịch vụ, thời gian; TTL 900 giây'],
    ['Chỉ mục slot', 'slot-index:{staffId}:{yyyy-mm-dd}', 'Sorted Set; score là thời gian bắt đầu; member là token'],
    ['Cache chi nhánh', 'cache:branches', 'Danh sách chi nhánh; TTL 300 giây'],
    ['Cache dịch vụ', 'cache:services:{branchId}', 'Dịch vụ theo chi nhánh; TTL 300 giây'],
], widths=[3.3, 6, 7.5])

paragraph('Lua script được sử dụng để kiểm tra giao thoa và tạo hold trong một thao tác nguyên tử. Cách này tránh tình trạng hai yêu cầu cùng đọc thấy slot trống rồi đồng thời ghi dữ liệu.')

doc.add_heading('5 Luồng đặt lịch và chống tranh chấp', level=1)
booking_steps = [
    'Khách chọn chi nhánh, dịch vụ, chuyên viên, ngày và giờ.',
    'Backend tính end_time từ start_time và tổng duration_minutes của các dịch vụ.',
    'Backend kiểm tra appointment đã xác nhận trong MongoDB.',
    'Redis Lua script kiểm tra các hold giao nhau và tạo khóa TTL 900 giây.',
    'Frontend nhận hold_token và hiển thị thời gian còn lại.',
    'Khách xác nhận; backend kiểm tra token, chủ sở hữu và dữ liệu đi kèm.',
    'MongoDB cập nhật staff_calendars nguyên tử rồi tạo appointment.',
    'Backend xóa hold khỏi Redis và trả mã đặt lịch.',
    'Nếu lỗi xảy ra, backend rollback slot vừa thêm trong MongoDB.'
]
for idx, step in enumerate(booking_steps, 1):
    paragraph(f'{idx}. {step}')

doc.add_heading('5 1 Điều kiện giao thoa thời gian', level=2)
paragraph('Hai khoảng thời gian giao nhau khi thời gian bắt đầu của lịch hiện có nhỏ hơn thời gian kết thúc mới và thời gian kết thúc của lịch hiện có lớn hơn thời gian bắt đầu mới.')
code('existing.start_time < new.end_time')
code('AND existing.end_time > new.start_time')

doc.add_heading('5 2 Mã phản hồi chính', level=2)
table(['HTTP', 'Trường hợp', 'Phản hồi mong đợi'], [
    ['201', 'Giữ chỗ hoặc đặt lịch thành công', 'Trả hold_token hoặc appointment'],
    ['400', 'Dữ liệu không hợp lệ', 'Thông báo trường hoặc nghiệp vụ sai'],
    ['401', 'Thiếu hoặc sai JWT', 'Yêu cầu đăng nhập lại'],
    ['403', 'Sai vai trò', 'Từ chối truy cập'],
    ['409', 'Slot đã bận, đang được giữ hoặc token hết hạn', 'Yêu cầu chọn khung giờ khác'],
], widths=[2.2, 7.1, 7.5])

doc.add_heading('6 Danh sách API', level=1)
doc.add_heading('6 1 Xác thực và danh mục', level=2)
table(['Phương thức', 'Endpoint', 'Chức năng', 'Quyền'], [
    ['POST', '/api/auth/register', 'Đăng ký khách hàng', 'Công khai'],
    ['POST', '/api/auth/login', 'Đăng nhập và nhận JWT', 'Công khai'],
    ['GET', '/api/auth/me', 'Thông tin tài khoản', 'Đã đăng nhập'],
    ['GET', '/api/branches', 'Danh sách hoặc tìm chi nhánh gần nhất', 'Công khai'],
    ['GET', '/api/services', 'Danh sách dịch vụ', 'Công khai'],
    ['GET', '/api/staff', 'Danh sách chuyên viên', 'Công khai'],
], widths=[2.5, 5, 6.8, 2.5])

doc.add_heading('6 2 Giữ chỗ và lịch hẹn', level=2)
table(['Phương thức', 'Endpoint', 'Chức năng', 'Quyền'], [
    ['POST', '/api/slots/hold', 'Giữ slot trong Redis 15 phút', 'CUSTOMER'],
    ['GET', '/api/slots/:token', 'Kiểm tra thời gian còn lại', 'CUSTOMER'],
    ['DELETE', '/api/slots/:token', 'Chủ động giải phóng slot', 'CUSTOMER'],
    ['POST', '/api/appointments', 'Xác nhận lịch từ hold_token', 'CUSTOMER'],
    ['GET', '/api/appointments/mine', 'Danh sách lịch theo người dùng', 'Đã đăng nhập'],
    ['PATCH', '/api/appointments/:id/status', 'Cập nhật trạng thái phục vụ', 'STAFF ADMIN'],
    ['PATCH', '/api/appointments/:id/cancel', 'Khách hủy lịch', 'CUSTOMER'],
], widths=[2.5, 5.2, 6.6, 2.5])

doc.add_heading('6 3 Đánh giá quản trị và báo cáo', level=2)
table(['Nhóm', 'Endpoint tiêu biểu', 'Kết quả'], [
    ['Đánh giá', 'POST /api/reviews; GET /api/reviews', 'Tạo và tra cứu đánh giá'],
    ['Quản trị', '/api/admin/branches; /services; /staff', 'CRUD danh mục và nhân viên'],
    ['Báo cáo', '/api/reports/dashboard', 'Doanh thu, top dịch vụ và hiệu suất nhân viên'],
], widths=[3.2, 7, 6.6])

doc.add_heading('7 Kế hoạch giao diện', level=1)
table(['Giai đoạn', 'Màn hình', 'Nội dung triển khai'], [
    ['Khách hàng', 'Trang chủ, dịch vụ, đặt lịch, lịch của tôi, đánh giá', 'Luồng đầy đủ từ chọn dịch vụ đến xác nhận; hiển thị countdown TTL'],
    ['Nhân viên', 'Dashboard, lịch ngày tuần, chi tiết lịch', 'Xem khách sắp đến và chuyển trạng thái phục vụ'],
    ['Quản trị', 'Dashboard, chi nhánh, dịch vụ, nhân viên, báo cáo', 'CRUD danh mục, quản lý tài khoản và biểu đồ tổng hợp'],
], widths=[3.1, 6.1, 7.6])

doc.add_heading('8 Kế hoạch triển khai theo giai đoạn', level=1)
table(['Tuần', 'Trọng tâm', 'Công việc', 'Sản phẩm bàn giao'], [
    ['1', 'Phân tích và thiết kế', 'Chuẩn hóa yêu cầu, use case, kiến trúc, data model, Redis key và API', 'Use case, sơ đồ kiến trúc, mô hình dữ liệu, backlog'],
    ['2', 'Hạ tầng và xác thực', 'Docker Compose, MongoDB, Redis, JWT, Roles Guard, validation và Swagger', 'Hai CSDL connected, ba vai trò đăng nhập đúng'],
    ['3', 'Catalog và đặt lịch', 'Chi nhánh, dịch vụ, nhân viên, geo query, Redis hold và atomic booking', 'Không thể giữ hoặc đặt hai slot giao nhau'],
    ['4', 'Giao diện khách hàng', 'Trang chủ, form đặt lịch, countdown, lịch sử, hủy lịch và đánh giá', 'Luồng khách hàng hoàn chỉnh trên trình duyệt'],
    ['5', 'Nhân viên và quản trị', 'Lịch nhân viên, trạng thái, CRUD danh mục và dashboard', 'Phân quyền đúng và báo cáo hoạt động'],
    ['6', 'Kiểm thử và báo cáo', 'Unit, integration, E2E, explain index, dữ liệu demo, tài liệu và slide', 'Bản chạy ổn định và bộ hồ sơ bảo vệ'],
], widths=[1.5, 3.3, 7, 5])

doc.add_heading('8 1 Mốc kiểm soát', level=2)
table(['Mốc', 'Điều kiện qua mốc'], [
    ['M1 Hạ tầng', 'MongoDB và Redis khởi động bằng Docker; health endpoint báo connected'],
    ['M2 Nghiệp vụ', 'Redis TTL và MongoDB atomic booking hoạt động đúng dưới tải đồng thời'],
    ['M3 Giao diện', 'Khách hoàn thành đặt lịch; nhân viên cập nhật trạng thái; admin xem báo cáo'],
    ['M4 Bàn giao', 'Build, test, seed, Swagger, README và kịch bản demo đầy đủ'],
], widths=[4, 12.8])

doc.add_heading('9 Kế hoạch kiểm thử', level=1)
table(['Cấp độ', 'Phạm vi', 'Ca kiểm thử trọng tâm'], [
    ['Unit', 'Hàm nghiệp vụ độc lập', 'Tính thời lượng, giao thoa, chuyển trạng thái, quyền và DTO'],
    ['Integration', 'NestJS với MongoDB và Redis', 'JWT, hold TTL, cache, xác nhận lịch, snapshot và rollback'],
    ['Concurrency', 'Nhiều yêu cầu đồng thời', 'Chỉ một khách giữ được cùng slot; các yêu cầu khác nhận 409'],
    ['End to end', 'Toàn bộ hệ thống', 'Đăng nhập, giữ slot, đặt lịch, phục vụ, đánh giá và báo cáo'],
], widths=[3.1, 5.2, 8.6])

doc.add_heading('9 1 Bộ ca kiểm thử tối thiểu', level=2)
for item in [
    'Token JWT hợp lệ, hết hạn, bị sửa và sai vai trò.',
    'Slot không giao nhau, chạm biên và giao nhau một phần hoặc toàn phần.',
    'Hai tài khoản gửi yêu cầu giữ slot đồng thời.',
    'Hold tồn tại, hết hạn, thuộc người khác hoặc không khớp dữ liệu đặt lịch.',
    'MongoDB thất bại sau khi calendar đã cập nhật và hệ thống rollback đúng.',
    'Cache miss, cache hit và cache hết hạn.',
    'Aggregation trả đúng doanh thu, top dịch vụ và rating.'
]: bullet(item)

doc.add_heading('10 Tiêu chí nghiệm thu', level=1)
table(['Nhóm', 'Tiêu chí', 'Minh chứng'], [
    ['NoSQL', 'MongoDB và Redis đều được sử dụng thực tế', 'Dữ liệu, log, health endpoint và thao tác demo'],
    ['MongoDB', 'Có embedded document, GeoJSON, index và Aggregation Pipeline', 'Schema, explain và kết quả truy vấn'],
    ['Redis', 'Hold TTL 900 giây, khóa tranh chấp và cache', 'TTL, key, Sorted Set và phản hồi 409'],
    ['Bảo mật', 'JWT, hash mật khẩu và phân quyền ba vai trò', 'API test và Swagger'],
    ['Nghiệp vụ', 'Không xảy ra double-booking', 'Concurrency test'],
    ['Giao diện', 'Responsive và hoàn tất các luồng chính', 'Demo trình duyệt'],
    ['Chất lượng', 'Build và test đạt; có README, seed và Docker', 'Kết quả command và repository'],
], widths=[2.8, 8.2, 5.9])

doc.add_heading('11 Rủi ro và biện pháp xử lý', level=1)
table(['Rủi ro', 'Tác động', 'Biện pháp'], [
    ['Redis không hoạt động', 'Không thể giữ slot hoặc đọc cache', 'Health check, Docker restart policy; không bỏ qua khóa khi Redis lỗi'],
    ['Race condition', 'Hai lịch trùng nhân viên', 'Redis Lua script và MongoDB atomic update hai lớp'],
    ['Dữ liệu snapshot sai', 'Sai giá hoặc thời lượng lịch hẹn', 'Chỉ lấy dữ liệu dịch vụ từ backend tại thời điểm xác nhận'],
    ['JWT hoặc phân quyền sai', 'Rò rỉ dữ liệu giữa vai trò', 'Guard tập trung và test từng endpoint'],
    ['Index thiếu', 'Truy vấn chậm khi dữ liệu tăng', 'Kiểm tra explain và index theo query thực tế'],
    ['Lỗi môi trường demo', 'Không thể trình bày', 'Docker Compose, seed idempotent và checklist chạy thử trước bảo vệ'],
], widths=[4, 5.2, 7.7])

doc.add_heading('12 Kịch bản trình bày đồ án', level=1)
demo_steps = [
    'Giới thiệu bài toán và lý do lựa chọn MongoDB cùng Redis.',
    'Trình bày collection MongoDB và cấu trúc khóa Redis.',
    'Khởi động hệ thống bằng Docker Compose và mở health endpoint.',
    'Đăng nhập bằng hai tài khoản khách hàng.',
    'Khách thứ nhất giữ slot trong Redis.',
    'Khách thứ hai giữ cùng slot và nhận phản hồi 409.',
    'Khách thứ nhất xác nhận và tạo appointment trong MongoDB.',
    'Kiểm tra hold đã được xóa khỏi Redis.',
    'Nhân viên chuyển lịch sang IN_PROGRESS rồi COMPLETED.',
    'Khách hàng gửi đánh giá.',
    'Quản trị viên mở báo cáo doanh thu và top dịch vụ.',
    'Kết luận bằng kết quả kiểm thử và hướng mở rộng.'
]
for idx, step in enumerate(demo_steps, 1):
    paragraph(f'{idx}. {step}')

doc.add_heading('13 Ưu tiên triển khai tiếp theo', level=1)
paragraph('Bốn hạng mục cần ưu tiên trước khi hoàn thiện đồ án là API quản trị, giao diện nhân viên và quản trị viên, đồng hồ đếm ngược TTL trên frontend và integration test chạy đồng thời với MongoDB cùng Redis. Sau khi hoàn thành các hạng mục này, nhóm chuyển sang tối ưu truy vấn, chuẩn hóa dữ liệu demo, hoàn thiện báo cáo và luyện kịch bản bảo vệ.')

# Footer page number
for sec in doc.sections:
    footer = sec.footer
    p = footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run('Trang ')
    r.font.name = 'Arial'; r.font.size = Pt(9); r.font.color.rgb = RGBColor(90, 90, 90)
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    p._p.append(fld)

doc.core_properties.title = 'Kế hoạch triển khai hệ thống quản lý lịch hẹn Spa Salon'
doc.core_properties.subject = 'Đồ án NoSQL sử dụng MongoDB và Redis'
doc.core_properties.author = 'Nhóm thực hiện đồ án'
doc.save(OUT)
print(OUT)
